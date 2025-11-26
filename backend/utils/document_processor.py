import os
from pathlib import Path
from typing import Dict, List, Any
import PyPDF2
from docx import Document
import pandas as pd
from PIL import Image
import config

class DocumentProcessor:
    """Utility class for processing various document types"""

    @staticmethod
    async def process_file(file_path: Path, file_type: str) -> Dict[str, Any]:
        """
        Process a file based on its type and extract content

        Args:
            file_path: Path to the file
            file_type: Type of file (pdf, docx, csv, txt, etc.)

        Returns:
            Dictionary containing extracted content and metadata
        """
        try:
            if file_type == 'pdf':
                return await DocumentProcessor.process_pdf(file_path)
            elif file_type in ['doc', 'docx']:
                return await DocumentProcessor.process_docx(file_path)
            elif file_type in ['csv', 'xlsx', 'xls']:
                return await DocumentProcessor.process_spreadsheet(file_path, file_type)
            elif file_type == 'txt':
                return await DocumentProcessor.process_text(file_path)
            elif file_type in ['jpg', 'jpeg', 'png', 'bmp']:
                return await DocumentProcessor.process_image(file_path)
            else:
                return {
                    "success": False,
                    "error": f"Unsupported file type: {file_type}"
                }
        except Exception as e:
            return {
                "success": False,
                "error": str(e)
            }

    @staticmethod
    async def process_pdf(file_path: Path) -> Dict[str, Any]:
        """Extract text and metadata from PDF file"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)

                # Extract text from all pages
                pages_content = []
                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    pages_content.append({
                        "page_number": page_num + 1,
                        "content": text,
                        "char_count": len(text)
                    })

                # Get PDF metadata
                metadata = pdf_reader.metadata or {}

                return {
                    "success": True,
                    "file_type": "pdf",
                    "num_pages": num_pages,
                    "pages": pages_content,
                    "full_text": "\n\n".join([p["content"] for p in pages_content]),
                    "metadata": {
                        "title": metadata.get('/Title', ''),
                        "author": metadata.get('/Author', ''),
                        "subject": metadata.get('/Subject', ''),
                        "creator": metadata.get('/Creator', ''),
                    }
                }
        except Exception as e:
            return {
                "success": False,
                "error": f"Error processing PDF: {str(e)}"
            }

    @staticmethod
    async def process_docx(file_path: Path) -> Dict[str, Any]:
        """Extract text and metadata from DOCX file"""
        try:
            doc = Document(file_path)

            # Extract paragraphs
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # Extract tables
            tables_data = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables_data.append(table_data)

            full_text = "\n\n".join(paragraphs)

            return {
                "success": True,
                "file_type": "docx",
                "num_paragraphs": len(paragraphs),
                "num_tables": len(tables_data),
                "paragraphs": paragraphs,
                "tables": tables_data,
                "full_text": full_text,
                "metadata": {
                    "char_count": len(full_text),
                    "word_count": len(full_text.split())
                }
            }
        except Exception as e:
            return {
                "success": False,
                "error": f"Error processing DOCX: {str(e)}"
            }

    @staticmethod
    async def process_spreadsheet(file_path: Path, file_type: str) -> Dict[str, Any]:
        """Extract data from spreadsheet files (CSV, XLSX, XLS)"""
        try:
            if file_type == 'csv':
                df = pd.read_csv(file_path)
            else:
                df = pd.read_excel(file_path)

            # Convert to dictionary
            data_dict = df.to_dict(orient='records')

            # Get basic statistics
            stats = {
                "num_rows": len(df),
                "num_columns": len(df.columns),
                "columns": list(df.columns),
                "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()}
            }

            # Get a preview (first 10 rows)
            preview = df.head(10).to_dict(orient='records')

            return {
                "success": True,
                "file_type": file_type,
                "data": data_dict,
                "preview": preview,
                "statistics": stats,
                "full_text": df.to_string()
            }
        except Exception as e:
            return {
                "success": False,
                "error": f"Error processing spreadsheet: {str(e)}"
            }

    @staticmethod
    async def process_text(file_path: Path) -> Dict[str, Any]:
        """Extract content from plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()

            lines = content.split('\n')

            return {
                "success": True,
                "file_type": "txt",
                "content": content,
                "num_lines": len(lines),
                "num_chars": len(content),
                "num_words": len(content.split())
            }
        except Exception as e:
            return {
                "success": False,
                "error": f"Error processing text file: {str(e)}"
            }

    @staticmethod
    async def process_image(file_path: Path) -> Dict[str, Any]:
        """Extract metadata from image file"""
        try:
            with Image.open(file_path) as img:
                return {
                    "success": True,
                    "file_type": "image",
                    "format": img.format,
                    "size": img.size,
                    "mode": img.mode,
                    "width": img.width,
                    "height": img.height,
                    "info": dict(img.info)
                }
        except Exception as e:
            return {
                "success": False,
                "error": f"Error processing image: {str(e)}"
            }

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[Dict[str, Any]]:
        """
        Split text into chunks with overlap

        Args:
            text: Text to split
            chunk_size: Maximum size of each chunk (in characters)
            overlap: Number of characters to overlap between chunks

        Returns:
            List of dictionaries containing chunks and metadata
        """
        chunks = []
        start = 0
        text_length = len(text)
        chunk_index = 0

        while start < text_length:
            end = start + chunk_size
            chunk_text = text[start:end]

            chunks.append({
                "chunk_index": chunk_index,
                "text": chunk_text,
                "start_pos": start,
                "end_pos": end,
                "length": len(chunk_text)
            })

            start = end - overlap
            chunk_index += 1

        return chunks
